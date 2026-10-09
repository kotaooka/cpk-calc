# Excel 入出力テスト用のファイルと期待値を作る（python tests/make_xlsx_fixtures.py、openpyxl が必要）
#   fixtures/made.xlsx         : openpyxl で作った、複数シート・日本語・空セル・各種の値を含むファイル
#   fixtures/*.expected.json   : openpyxl で読んだ各シートの空でないセル [行, 列, 値]（0 始まり）
import json
from pathlib import Path
import openpyxl
D = Path(__file__).parent / 'fixtures'
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '測定データ'
ws.append(['No', '外径 φ10', '長さ & 幅 <mm>', None, '備考'])
for i in range(1, 9):
    ws.append([i, 10 + i * 0.001, 1e-7 * i if i % 2 else 123456789.125, None, '要確認' if i == 3 else None])
ws['G12'] = 'ぽつんと'; ws['B20'] = -0.5
ws2 = wb.create_sheet('二枚目'); ws2['A1'] = '行方向'; ws2['B1'] = 1.5; ws2['C1'] = True; ws2['A2'] = ' 前後に空白 '
wb.save(D / 'made.xlsx')
def dump(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None or v == '': continue
                if isinstance(v, bool): v = 'TRUE' if v else 'FALSE'
                cells.append([c.row - 1, c.column - 1, v])
        out.append({'name': ws.title, 'cells': cells})
    return out
for name in ['made.xlsx', '../../samples/sample.xlsx']:
    p = (D / name).resolve()
    (D / (p.stem + '.expected.json')).write_text(json.dumps(dump(p), ensure_ascii=False), encoding='utf-8', newline='\n')
print('作成しました')
